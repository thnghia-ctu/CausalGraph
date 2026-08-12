from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.concept_builder.concept_clustering import MIN_SENTENCE_COUNT

DETAIL_COLUMNS = [
    "source_direction", "subject_text", "predicate", "object_text",
    "target_direction", "original_sentence",
]

DIRECTION_LABELS = {1: "TĂNG", -1: "GIẢM"}

REPORT_COLUMNS = [
    "edge_id", "Quan hệ",
    "Đồng ý", "Không đồng ý", "Lưỡng lự", "Ý kiến",
]

REVIEW_COLUMN_FILLS = {
    "Đồng ý": "C6EFCE",
    "Không đồng ý": "FFC7CE",
    "Lưỡng lự": "FFEB9C",
}
EMPTY_REVIEW = dict.fromkeys(REVIEW_COLUMN_FILLS | {"Ý kiến": ""}, "")

COLUMN_WIDTHS = {
    "edge_id": 8, "Quan hệ": 90,
    "Đồng ý": 9, "Không đồng ý": 12, "Lưỡng lự": 9, "Ý kiến": 32,
}

HEADER_ROW_HEIGHT = 22
DETAIL_ROW_HEIGHT = 75
BLOCK_TOP_BORDER = Border(top=Side(style="thick"))
BLOCK_BORDER_COLUMNS = ("edge_id", "Quan hệ")


def _column_index(name: str) -> int:
    return REPORT_COLUMNS.index(name) + 1


def _format_direction(direction) -> str:
    try:
        return DIRECTION_LABELS.get(int(direction), "")
    except (TypeError, ValueError):
        return ""


def format_edge_relation(source_label: str, target_label: str) -> str:
    return f"{source_label} ------> {target_label}"


def format_component_relation(detail: dict) -> CellRichText:
    source_direction = _format_direction(detail.get("source_direction"))
    target_direction = _format_direction(detail.get("target_direction"))
    source_prefix = f"({source_direction}) " if source_direction else ""
    target_suffix = f" ({target_direction})" if target_direction else ""
    relation_line = (
        f"{source_prefix}{detail.get('subject_text')} "
        f"---{detail.get('predicate')}---> "
        f"{detail.get('object_text')}{target_suffix}"
    )
    original_sentence = str(detail.get("original_sentence") or "")
    return CellRichText(
        f"{relation_line}\n\n",
        TextBlock(InlineFont(i=True), f"(Câu gốc: {original_sentence})"),
    )


def _row(edge_id: str, relation, is_header: bool) -> dict:
    return {"edge_id": edge_id, "Quan hệ": relation, **EMPTY_REVIEW, "_is_header": is_header}


def build_edge_rows(
    relations: pd.DataFrame,
    concepts: pd.DataFrame,
    *,
    min_sentence_count: int = MIN_SENTENCE_COUNT,
) -> list[dict]:
    labels = dict(zip(concepts["concept_id"], concepts["representative_label"]))

    relations = relations.dropna(subset=["source_concept_id", "target_concept_id"])
    grouped = relations.groupby(["source_concept_id", "target_concept_id"])
    sentence_counts = grouped["simple_sentence"].nunique()
    edge_keys = sentence_counts[sentence_counts >= min_sentence_count].sort_values(
        ascending=False,
    ).index

    rows = []
    for rank, (source_id, target_id) in enumerate(edge_keys, start=1):
        edge_id = f"E{rank:04d}"
        source_label = labels.get(source_id, source_id)
        target_label = labels.get(target_id, target_id)

        rows.append(_row(edge_id, format_edge_relation(source_label, target_label), True))
        for _, detail in grouped.get_group((source_id, target_id)).iterrows():
            detail = {col: detail.get(col) for col in DETAIL_COLUMNS}
            rows.append(_row(edge_id, format_component_relation(detail), False))
    return rows


def export_report(rows: list[dict], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Đánh giá quan hệ"
    sheet.sheet_properties.outlinePr.summaryBelow = False

    sheet.append(REPORT_COLUMNS)
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    relation_column = _column_index("Quan hệ")
    for row in rows:
        sheet.append([row[column] for column in REPORT_COLUMNS])
        excel_row = sheet.max_row

        if row["_is_header"]:
            sheet.row_dimensions[excel_row].height = HEADER_ROW_HEIGHT
            for cell in sheet[excel_row]:
                cell.font = Font(bold=True)
            for column_name in BLOCK_BORDER_COLUMNS:
                sheet.cell(row=excel_row, column=_column_index(column_name)).border = BLOCK_TOP_BORDER
        else:
            sheet.row_dimensions[excel_row].height = DETAIL_ROW_HEIGHT
            sheet.row_dimensions[excel_row].outline_level = 1
            sheet.cell(row=excel_row, column=relation_column).alignment = Alignment(
                wrap_text=True, vertical="top",
            )

    for column_name, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[get_column_letter(_column_index(column_name))].width = width

    last_row = sheet.max_row
    for column_name, color in REVIEW_COLUMN_FILLS.items():
        column_letter = get_column_letter(_column_index(column_name))
        cell_range = f"{column_letter}2:{column_letter}{last_row}"
        sheet.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[f"NOT(ISBLANK({column_letter}2))"],
                fill=PatternFill("solid", fgColor=color),
            ),
        )

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_COLUMNS))}{last_row}"

    workbook.save(output_path)
    return output_path


def build_report(
    relations: pd.DataFrame,
    concepts: pd.DataFrame,
    output_path: str | Path,
    *,
    min_sentence_count: int = MIN_SENTENCE_COUNT,
) -> Path:
    rows = build_edge_rows(relations, concepts, min_sentence_count=min_sentence_count)
    return export_report(rows, output_path)
